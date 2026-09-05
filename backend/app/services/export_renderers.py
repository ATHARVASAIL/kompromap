"""DOCX and PDF export renderers.

DOCX uses python-docx and ships a styled document that any word processor
can open and the client can mark up with tracked changes. It mirrors the
Markdown structure (executive summary, findings, chains, remediation) but
uses proper heading styles and a severity table.

PDF uses WeasyPrint on an HTML intermediate — it's the most reliable
cross-platform approach (wkhtmltopdf is unmaintained, reportlab requires
hand-laying every element). The HTML is self-contained: all styles are
inline or in a <style> block, no external assets. If WeasyPrint's system
dependencies (libpango, libcairo) aren't available, the endpoint falls
back gracefully to returning the HTML with a note.
"""
from __future__ import annotations

import io
import logging
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.engagement_report import EngagementReport

logger = logging.getLogger(__name__)

# ── colour palette (matches the dark-theme UI) ──────────────────────────
SEV_HEX = {
    "Critical": "C0392B",
    "High": "E67E22",
    "Medium": "F1C40F",
    "Low": "3498DB",
    "Informational": "95A5A6",
}


# =====================================================================
# DOCX
# =====================================================================
def render_docx(report: EngagementReport) -> bytes:
    doc = Document()

    # --- default body style ---
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    h = doc.add_heading(level=0)
    title = f"Penetration Test Report — {report.engagement_name}"
    if report.client_name:
        title += f" — {report.client_name}"
    run = h.add_run(title)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0A, 0x0D, 0x12)

    sub = doc.add_paragraph()
    sub.add_run(
        f"Generated {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')} by Kompromap"
    ).italic = True
    sub.runs[0].font.color.rgb = RGBColor(0x90, 0x99, 0xAC)

    doc.add_paragraph("")  # spacer

    # --- Executive summary ---
    doc.add_heading("Executive summary", level=1)
    s = report.summary
    counts = s["severity_counts"]
    doc.add_paragraph(
        f"This engagement identified {s['total_findings']} findings across "
        f"{s['total_nodes']} mapped assets and services."
    )
    if s["chain_count"]:
        doc.add_paragraph(
            f"More importantly, {s['chain_count']} viable attack chains were identified "
            f"linking an internet-facing entry point to a business-critical asset."
        )

    # severity table
    _add_severity_table(doc, counts)

    # --- Attack chains ---
    doc.add_heading("Attack chains", level=1)
    if not report.chains:
        doc.add_paragraph(
            "No complete attack chain from an entry point to a crown jewel was identified. "
            "See Caveats and limitations for what that does and does not mean."
        )
    for chain in report.chains:
        doc.add_heading(
            f"Chain #{chain.rank}: {chain.entry_point} → {chain.crown_jewel}", level=2
        )
        doc.add_paragraph(
            f"Cost: {chain.total_cost}  |  Steps: {len(chain.steps)}  |  "
            f"Exploitation steps: {chain.exploit_step_count}"
        )
        if chain.narrative:
            doc.add_paragraph(chain.narrative)

        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Light Grid Accent 1"
        hdr = tbl.rows[0].cells
        for i, label in enumerate(["#", "From", "Relationship", "To"]):
            hdr[i].text = label
            hdr[i].paragraphs[0].runs[0].bold = True
        for st in chain.steps:
            row = tbl.add_row().cells
            row[0].text = str(st["index"])
            row[1].text = st["from"]
            row[2].text = st["relationship"]
            row[3].text = st["to"]
        doc.add_paragraph("")

    # --- Findings ---
    doc.add_heading("Findings", level=1)
    for f in report.findings:
        h = doc.add_heading(f.title, level=2)
        try:
            hex_color = SEV_HEX.get(f.severity, "000000")
            h.runs[0].font.color.rgb = RGBColor(
                int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
            )
        except Exception:
            pass

        meta_parts = [f"Severity: {f.severity}"]
        if f.cvss_score is not None:
            meta_parts.append(f"CVSS: {f.cvss_score}")
        if f.cwe:
            meta_parts.append(f"CWE: {f.cwe}")
        if f.owasp_category:
            meta_parts.append(f"OWASP: {f.owasp_category}")
        meta_parts.append(f"Status: {f.status}")
        doc.add_paragraph("  |  ".join(meta_parts))

        if f.cvss_vector:
            p = doc.add_paragraph()
            r = p.add_run(f.cvss_vector)
            r.font.name = "Courier New"
            r.font.size = Pt(9)

        if f.affected:
            doc.add_paragraph(f"Affected: {', '.join(f.affected)}")

        doc.add_paragraph(f"Evidence: {f.evidence}" if f.evidence else "No evidence recorded.")

        flags = []
        if f.exploit_public:
            flags.append("Public exploit available")
        if not f.auth_required:
            flags.append("No authentication required")
        if f.in_chain:
            flags.append("Appears on an attack chain")
        if not f.complexity_measured:
            flags.append("Complexity assumed (no CVSS vector)")
        if flags:
            p = doc.add_paragraph()
            p.add_run("; ".join(flags)).italic = True
        doc.add_paragraph("")

    # --- Remediation ---
    doc.add_heading("Remediation priority", level=1)
    doc.add_paragraph(
        "Ranked by chain impact rather than raw severity. "
        "A medium finding on the cheapest path to a crown jewel matters more than an unreachable critical."
    )
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    for i, label in enumerate(["#", "Finding", "Severity", "Why"]):
        hdr[i].text = label
        hdr[i].paragraphs[0].runs[0].bold = True
    for item in report.remediation[:25]:
        row = tbl.add_row().cells
        row[0].text = str(item["rank"])
        row[1].text = item["title"]
        row[2].text = item["severity"]
        row[3].text = "; ".join(item["rationale"])

    # --- Scope ---
    doc.add_heading("Scope", level=1)
    for label, key in [
        ("Entry points", "entry_points"),
        ("Crown jewels", "crown_jewels"),
        ("Assets", "assets"),
        ("Services", "services"),
        ("Web applications", "web_applications"),
        ("Endpoints", "endpoints"),
        ("Data stores", "data_stores"),
    ]:
        values = report.scope.get(key) or []
        if values:
            shown = ", ".join(values[:60])
            doc.add_paragraph(f"{label} ({len(values)}): {shown}")

    # --- Caveats ---
    if report.caveats:
        doc.add_heading("Caveats and limitations", level=1)
        for c in report.caveats:
            doc.add_paragraph(c, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_severity_table(doc: Document, counts: dict):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Severity"
    hdr[1].text = "Count"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True
    for sev in ["Critical", "High", "Medium", "Low", "Informational"]:
        row = tbl.add_row().cells
        row[0].text = sev
        row[1].text = str(counts.get(sev, 0))
    doc.add_paragraph("")


# =====================================================================
# PDF (via WeasyPrint HTML → PDF)
# =====================================================================
def render_pdf(report: EngagementReport) -> tuple[bytes, str]:
    """Return (bytes, format).

    Tries WeasyPrint; if the system dependencies are missing (common
    in slim containers) returns (html_bytes, 'html-fallback') so the
    caller can still deliver something useful.
    """
    html_content = _pdf_html(report)
    try:
        from weasyprint import HTML  # lazy: not everyone has the deps

        pdf_io = io.BytesIO()
        HTML(string=html_content, base_url=".").write_pdf(pdf_io)
        return pdf_io.getvalue(), "pdf"
    except (ImportError, OSError) as exc:
        logger.warning("WeasyPrint unavailable (%s); returning HTML fallback", exc)
        return html_content.encode(), "html-fallback"


def _pdf_html(report: EngagementReport) -> str:
    e = __import__("html").escape
    s = report.summary
    counts = s["severity_counts"]

    title = e(report.engagement_name)
    if report.client_name:
        title += f" — {e(report.client_name)}"

    def sev_pill(sev: str) -> str:
        c = SEV_HEX.get(sev, "6B7488")
        return (
            f'<span class="pill" style="color:#{c};border-color:#{c}55;'
            f'background:#{c}18">{e(sev)}</span>'
        )

    parts: list[str] = []
    add = parts.append

    add("""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>Penetration Test Report</title>
<style>
:root{--bg:#fff;--panel:#f4f6f8;--line:#d0d5dd;--text:#11151c;
      --muted:#4a5261;--dim:#6b7280;--accent:#0ea5e9}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);
    font:11pt/1.65 ui-sans-serif,system-ui,sans-serif}
.wrap{max-width:940px;margin:0 auto;padding:36px 28px 72px}
h1{font-size:26px;margin:0 0 4px;letter-spacing:-.02em}
h2{font-size:17px;margin:40px 0 12px;padding-bottom:6px;
    border-bottom:1px solid var(--line);page-break-after:avoid}
h3{font-size:14px;margin:24px 0 8px;page-break-after:avoid}
h4{font-size:12px;margin:16px 0 4px}
.sub{color:var(--dim);font-size:12px;margin-bottom:20px}
table{width:100%;border-collapse:collapse;margin:10px 0;font-size:10.5pt}
th,td{text-align:left;padding:6px 8px;border-bottom:1px solid var(--line);
       vertical-align:top}
th{color:var(--muted);font-weight:600;font-size:10px;
    text-transform:uppercase;letter-spacing:.04em}
code,.mono{font-family:ui-monospace,Consolas,monospace;font-size:10.5pt}
code{background:var(--panel);border:1px solid var(--line);
     border-radius:3px;padding:1px 4px;color:var(--muted)}
.pill{display:inline-block;border:1px solid;border-radius:99px;
      padding:1px 8px;font-size:10px;font-weight:600}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(130px,1fr));
        gap:10px;margin:16px 0}
.card{background:var(--panel);border:1px solid var(--line);
      border-radius:8px;padding:12px 14px}
.card .n{font-size:22px;font-weight:600}
.card .l{color:var(--dim);font-size:10px;text-transform:uppercase;
         letter-spacing:.04em}
.bar{display:flex;height:6px;border-radius:99px;overflow:hidden;
     background:var(--panel);margin:6px 0 16px}
.finding{border-left:3px solid var(--line);padding-left:12px;margin:18px 0}
.flags{color:var(--dim);font-size:11px;font-style:italic}
.note{background:var(--panel);border:1px solid var(--line);
      border-left:3px solid var(--accent);border-radius:4px;
      padding:10px 14px;margin:12px 0;font-size:12px;color:var(--muted)}
footer{margin-top:40px;padding-top:14px;border-top:1px solid var(--line);
       color:var(--dim);font-size:10px}
.chain{background:var(--panel);border:1px solid var(--line);
        border-radius:8px;padding:14px 16px;margin:14px 0;
        page-break-inside:avoid}
@media print{
  body{font-size:10pt}
  .wrap{max-width:none;padding:0}
  h2{page-break-after:avoid}
  .chain,.finding{page-break-inside:avoid}
}
</style></head><body><div class="wrap">""")

    add(f"<h1>Penetration Test Report</h1>")
    add(f'<div class="sub">{title} · generated '
        f'{report.generated_at.strftime("%Y-%m-%d %H:%M UTC")} by Kompromap</div>')

    # Summary
    add("<h2>Executive summary</h2>")
    add('<div class="cards">')
    for label, val, color in [
        ("Findings", s["total_findings"], None),
        ("Attack chains", s["chain_count"], "#C0392B" if s["chain_count"] else None),
        ("Entry points", s["entry_point_count"], "#0ea5e9"),
        ("Crown jewels", s["crown_jewel_count"], "#C0392B"),
        ("Mapped nodes", s["total_nodes"], None),
    ]:
        style = f' style="color:{color}"' if color else ""
        add(f'<div class="card"><div class="l">{e(label)}</div>'
            f'<div class="n"{style}>{val}</div></div>')
    add("</div>")

    add("<p>" + e(s["total_findings"]) + " findings across " + e(str(s["total_nodes"])) + " nodes.</p>")

    # severity bar
    total = max(sum(counts.values()), 1)
    add('<div class="bar">')
    for sev in ["Critical", "High", "Medium", "Low", "Informational"]:
        n = counts.get(sev, 0)
        if n:
            c = SEV_HEX.get(sev, "6B7488")
            pct = n / total * 100
            add(f'<div style="width:{pct}%;background:#{c}"></div>')
    add("</div>")

    add("<table><tr><th>Severity</th><th>Count</th></tr>")
    for sev in ["Critical", "High", "Medium", "Low", "Informational"]:
        add(f"<tr><td>{sev_pill(sev)}</td><td>{counts.get(sev, 0)}</td></tr>")
    add("</table>")

    # Attack chains
    add("<h2>Attack chains</h2>")
    if not report.chains:
        add("<p><em>No chains were computed. See caveats.</em></p>")
    for chain in report.chains:
        add(f"<div class=\"chain\"><h3>Chain #{chain.rank}: {e(chain.entry_point)} → {e(chain.crown_jewel)}</h3>")
        add(f"<p><strong>Cost:</strong> {chain.total_cost} &nbsp;|&nbsp; "
            f"<strong>Steps:</strong> {len(chain.steps)} &nbsp;|&nbsp; "
            f"<strong>Exploitation steps:</strong> {chain.exploit_step_count}</p>")
        if chain.narrative:
            add(f"<p>{e(chain.narrative)}</p>")
        add("<table><tr><th>#</th><th>From</th><th>Relationship</th><th>To</th></tr>")
        for st in chain.steps:
            marker = " <em>(exploit)</em>" if st["is_exploit"] else ""
            add(f"<tr><td>{st['index']}</td><td>{e(st['from'])}</td>"
                f"<td>{e(st['relationship'])}{marker}</td>"
                f"<td>{e(st['to'])}</td></tr>")
        add("</table></div>")

    # Findings
    add("<h2>Findings</h2>")
    for f in report.findings:
        add(f"<div class=\"finding\"><h4>{e(f.title)}</h4>")
        meta = [f"<strong>Severity:</strong> {sev_pill(f.severity)}"]
        if f.cvss_score is not None:
            meta.append(f"<strong>CVSS:</strong> {f.cvss_score}")
        if f.cwe:
            meta.append(f"<strong>CWE:</strong> {f.cwe}")
        if f.owasp_category:
            meta.append(f"<strong>OWASP:</strong> {f.owasp_category}")
        meta.append(f"<strong>Status:</strong> {f.status}")
        add("<p>" + " &nbsp;|&nbsp; ".join(meta) + "</p>")
        if f.cvss_vector:
            add(f"<p><code>{e(f.cvss_vector)}</code></p>")
        if f.affected:
            add(f"<p><strong>Affected:</strong> {', '.join(e(a) for a in f.affected)}</p>")
        add(f"<p><strong>Evidence:</strong> {e(f.evidence) if f.evidence else '_No evidence recorded._'}</p>")
        flags = []
        if f.exploit_public:
            flags.append("Public exploit available")
        if not f.auth_required:
            flags.append("No authentication required")
        if f.in_chain:
            flags.append("Appears on an attack chain")
        if not f.complexity_measured:
            flags.append("Complexity assumed (no CVSS vector)")
        if flags:
            add(f"<p class=\"flags\">{'; '.join(flags)}</p>")
        add("</div>")

    # Remediation
    add("<h2>Remediation priority</h2>")
    add("<div class=\"note\">Ranked by chain impact rather than raw severity — a medium finding "
        "on the cheapest path to a crown jewel matters more than an unreachable critical.</div>")
    add("<table><tr><th>#</th><th>Finding</th><th>Severity</th><th>Why</th>"
        "<th>Breaks a chain</th></tr>")
    for item in report.remediation[:25]:
        add(f"<tr><td>{item['rank']}</td><td>{e(item['title'])}</td>"
            f"<td>{sev_pill(item['severity'])}</td>"
            f"<td class=\"flags\">{e('; '.join(item['rationale']))}</td>"
            f"<td>{'yes' if item['breaks_chain'] else 'no'}</td></tr>")
    add("</table>")

    # Scope
    add("<h2>Scope</h2><table>")
    for label, key in [
        ("Entry points", "entry_points"), ("Crown jewels", "crown_jewels"),
        ("Assets", "assets"), ("Services", "services"),
        ("Web applications", "web_applications"), ("Endpoints", "endpoints"),
        ("Data stores", "data_stores"),
    ]:
        vals = report.scope.get(key) or []
        if vals:
            shown = ", ".join(vals[:60])
            more = f" <span class=\"flags\">…and {len(vals) - 60} more</span>" if len(vals) > 60 else ""
            add(f"<tr><th>{e(label)} ({len(vals)})</th><td>{e(shown)}{more}</td></tr>")
    add("</table>")

    # Caveats
    if report.caveats:
        add("<h2>Caveats and limitations</h2><ul>")
        for c in report.caveats:
            add(f"<li>{e(c)}</li>")
        add("</ul>")

    add("<footer>Scoring model: ease = 0.4·CVSS + 0.3·public exploit + 0.2·no-auth + "
        "0.1·(1−complexity), with complexity read from each finding's CVSS vector where "
        "available. Path cost = sum of (1 − ease) across exploitation steps; lower cost "
        "means an easier attack.</footer>")
    add("</div></body></html>")
    return "\n".join(parts)
